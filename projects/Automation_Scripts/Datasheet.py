from openpyxl import load_workbook
from datetime import datetime
import pandas as pd
import glob
import os

# ==========================
# INPUT
# ==========================

dashboard_date_str = input(
    "Enter Dashboard Date (DD-MMM-YYYY): "
)

dashboard_date = datetime.strptime(
    dashboard_date_str,
    "%d-%b-%Y"
)

SOURCE_FILE = "Case Validation.xlsx"

MASTER_TRACKER = "Infy_Conventional-Master Tracker.xlsx"

files = glob.glob(
    "Infosys_Conventional-Datasheet*.xlsx"
)

if not files:
    raise FileNotFoundError(
        "Datasheet file not found"
    )

TARGET_FILE = max(
    files,
    key=os.path.getmtime
)

print(f"Target File : {TARGET_FILE}")

VALID_COLORS = {
    "GREEN",
    "RED",
    "ORANGE",
    "MAROON"
}

# ==========================
# LOAD FILES
# ==========================

dst_wb = load_workbook(TARGET_FILE)

dst_ws = dst_wb["Data"]

tracker_pj = pd.read_excel(
    MASTER_TRACKER,
    sheet_name="Pre-Joining"
)

# ==========================
# HELPERS
# ==========================

def get_next_row(ws):

    row = 3

    while ws[f"B{row}"].value:
        row += 1

    return row


def to_dt(value):

    if value is None:
        return None

    if isinstance(value, datetime):
        return value

    try:
        return pd.to_datetime(value)
    except:
        return None


def valid_result(value):

    if value is None:
        return "NA"

    value = str(value).strip().upper()

    if value in VALID_COLORS:
        return value

    return "NA"


def old_month_date(check_date, dashboard_date):

    if not check_date:
        return False

    return (
        check_date.year < dashboard_date.year
        or
        (
            check_date.year == dashboard_date.year
            and
            check_date.month < dashboard_date.month
        )
    )

def find_existing_row(ws, cid, case_start=None):

    target_case_start = pd.to_datetime(
        case_start,
        errors="coerce"
    )

    target_case_date = (
        target_case_start.date()
        if pd.notna(target_case_start)
        else None
    )

    for r in range(2, ws.max_row + 1):

        if (
            str(ws[f"B{r}"].value).strip()
            ==
            str(cid).strip()
        ):

            # A candidate can have multiple distinct cases/add-ons. Reuse a
            # row only for the same case start date; a later add-on gets its
            # own row and leaves the earlier case untouched.
            if target_case_date is None:
                return r

            existing_case_start = pd.to_datetime(
                ws[f"G{r}"].value,
                errors="coerce"
            )

            if (
                pd.notna(existing_case_start)
                and existing_case_start.date() == target_case_date
            ):
                return r

    return None

def update_if_blank(cell, value):

    if (
        cell.value is None
        or str(cell.value).strip() == ""
        or str(cell.value).strip().upper() == "NA"
    ):
        cell.value = value

def update_date_result_cells(row_num, record, mappings):

    for col, source in mappings.items():

        value = record.get(source)

        if pd.notna(value):

            dst_ws[f"{col}{row_num}"] = value

def build_gdc_key(record):

    return (
        str(record.get("GDB_SDate", "")).strip()
        + "|"
        + str(record.get("GDB_EDate", "")).strip()
        + "|"
        + str(record.get("GDB_Result", "")).strip().upper()
    )


def is_gdc_duplicate(rec, tracker_df):

    cid = str(rec.get("C_ID")).strip()

    current_key = build_gdc_key(rec)
    
    previous_rows = tracker_df[
        tracker_df["C_ID"].astype(str).str.strip()
        == cid
    ]

    delivery_date = None

    delivery_cols = [
        "FR-CAN ID_Date",
        "SR-CAN ID_Date",
        "SR2-CAN ID_Date",
        "SR3-CAN ID_Date",
        "FR-EMP ID_Date",
        "SR-EMP ID_Date",
        "SR2-EMP ID_Date",
        "SR3-EMP ID_Date"
    ]

    for col in delivery_cols:


        dt = pd.to_datetime(
            rec.get(col),
            errors="coerce"
        )

        if pd.notna(dt):
            delivery_date = dt
            break

    for _, prev in previous_rows.iterrows():

        old_key = build_gdc_key(prev)
        
        prev_date = None

        for col in delivery_cols:
        
            dt = pd.to_datetime(
                prev.get(col),
                errors="coerce"
            )

            if pd.notna(dt):
                prev_date = dt
                break

        if (
            old_key == current_key
            and pd.notna(prev_date)
            and pd.notna(delivery_date)
            and prev_date < delivery_date
        ):
            return True

    return False

def build_check_key(
    record,
    s_col,
    r_col,
    e_col
):

    return (
        str(record.get(s_col, "")).strip()
        + "|"
        + str(record.get(e_col, "")).strip()
        + "|"
        + str(record.get(r_col, "")).strip().upper()
    )


def is_duplicate_check(
    rec,
    tracker_df,
    s_col,
    r_col,
    e_col
):

    cid = str(
        rec.get("C_ID")
    ).strip()

    rows = tracker_df[
        tracker_df["C_ID"].astype(str).str.strip()
        == cid
    ]

    if len(rows) <= 1:
        return False

    current_row = None

    for _, row_data in rows.iterrows():

        for col in [
            "FR-CAN ID_Date",
            "SR-CAN ID_Date",
            "SR2-CAN ID_Date",
            "SR3-CAN ID_Date",
            "FR-EMP ID_Date",
            "SR-EMP ID_Date",
            "SR2-EMP ID_Date",
            "SR3-EMP ID_Date"
        ]:

            dt = pd.to_datetime(
                row_data.get(col),
                errors="coerce"
            )

            if (
                pd.notna(dt)
                and dt.date() == dashboard_date.date()
            ):
                current_row = row_data
                break

        if current_row is not None:
            break

    if current_row is None:
       
        return False

    current_key = build_check_key(
        current_row,
        s_col,
        r_col,
        e_col
    )

    for _, prev in rows.iloc[:-1].iterrows():

        old_key = build_check_key(
            prev,
            s_col,
            r_col,
            e_col
        )

        if (
            current_key == old_key
            and current_key != "||"
        ):
            return True

    return False

def should_count_reint(rec, tracker_df):

    cid = str(
        rec.get("C_ID")
    ).strip()

    rows = tracker_df[
        tracker_df["C_ID"].astype(str).str.strip()
        == cid
    ]

    if len(rows) == 0:
        return False

    tracker_row = rows.iloc[0]

    delivery_cols = [
        "FR-CAN ID_Date",
        "SR-CAN ID_Date",
        "SR2-CAN ID_Date",
        "SR3-CAN ID_Date",
        "FR-EMP ID_Date",
        "SR-EMP ID_Date",
        "SR2-EMP ID_Date",
        "SR3-EMP ID_Date"
    ]

    current_delivery = None

    for col in delivery_cols:

        dt = pd.to_datetime(
            tracker_row.get(col),
            errors="coerce"
        )

        if (
            pd.notna(dt)
            and dt.date() == dashboard_date.date()
        ):
            current_delivery = dt
            break

    if current_delivery is None:
        return False

    previous_delivery = None

    for col in delivery_cols:

        dt = pd.to_datetime(
            tracker_row.get(col),
            errors="coerce"
        )

        if (
            pd.notna(dt)
            and dt < current_delivery
        ):

            if (
                previous_delivery is None
                or dt > previous_delivery
            ):
                previous_delivery = dt

    reint_edate = pd.to_datetime(
        tracker_row.get("ReINT_GDB_EDate"),
        errors="coerce"
    )

    if (
        pd.notna(reint_edate)
        and previous_delivery is not None
        and reint_edate > previous_delivery
    ):
        return True

    return False

# ==========================
# DATASHEET COLUMN MAP
# ==========================

COMMON_MAP = {
    "C_ID": "B",
    "Emp_No": "C",
    "C_Name": "D",
    "Package Name": "E",
    "Checks to be Initiate": "F",
    "Case_SDate": "G",
    "Report Type": "H"
}


# ==========================
# PJ CHECK MAPPING
# ==========================

PJ_MAP = {
    "Ed_Result": "Y",
    "CE1_Result": "Z",
    "E2_Result": "AA",
    "E3_Result": "AB",
    "E4_Result": "AC",
    "E5_Result": "AD",
    "E6_Result": "AE",
    "eP_CRC_Result": "AF",
    "GDB_Result": "AG"
}

PE_MAP = {
    "Ed_Result": "Y",
    "CE1_Result": "Z",
    "E2_Result": "AA",
    "E3_Result": "AB",
    "E4_Result": "AC",
    "E5_Result": "AD",
    "eP_CRC_Result": "AF",
    "GDB_Result": "AG"
}

FRESHER_MAP = {
    "Ed_Result": "Y",
    "eP_CRC_Result": "AF",
    "GDB_Result": "AG"
}

PJ_DATE_MAPPING = {
    "I": "FR-CAN ID_Date",
    "J": "FR-CAN ID_Result",
    "K": "SR-CAN ID_Date",
    "L": "SR-CAN ID_Result",
    "M": "SR2-CAN ID_Date",
    "N": "SR2-CAN ID_Result",
    "O": "SR3-CAN ID_Date",
    "P": "SR3-CAN ID_Result",
    "Q": "FR-EMP ID_Date",
    "R": "FR-EMP ID_Result",
    "S": "SR-EMP ID_Date",
    "T": "SR-EMP ID_Result",
    "U": "SR2-EMP ID_Date",
    "V": "SR2-EMP ID_Result",
    "W": "SR3-EMP ID_Date",
    "X": "SR3-EMP ID_Result"
}

PE_DATE_MAPPING = {
    "I": "IR_Date",
    "J": "IR_Result",

    "Q": "FR_Date",
    "R": "FR_Result",

    "S": "SR_Date",
    "T": "SR_Result",

    "U": "SR2_Date",
    "V": "SR2_Result",

    "W": "SR3_Date",
    "X": "SR3_Result"
}

# ==========================
# WRITE COMMON
# ==========================

def write_common(row_num, record):

    cid = record.get("C_ID")

    if pd.notna(cid) and str(cid).strip() != "":
        dst_ws[f"B{row_num}"] = cid
    else:
        dst_ws[f"B{row_num}"] = "NA"

    emp_no = record.get("Emp_No")

    if (
        pd.notna(emp_no)
        and str(emp_no).strip() != ""
        and str(emp_no).strip().upper() != "NA"
    ):
        dst_ws[f"C{row_num}"] = emp_no
    else:
        dst_ws[f"C{row_num}"] = "NA"

    dst_ws[f"D{row_num}"] = record["C_Name"]
    dst_ws[f"E{row_num}"] = record["Package Name"]
    dst_ws[f"F{row_num}"] = record["Checks to be Initiate"]
    dst_ws[f"G{row_num}"] = record["Case_SDate"]
    dst_ws[f"H{row_num}"] = record["Report Type"]

    additional_cost = record.get("Additional Cost")

    if (
        pd.notna(additional_cost)
        and str(additional_cost).strip() != ""
        and str(additional_cost).strip().upper() != "NA"
    ):
        dst_ws[f"AH{row_num}"] = additional_cost

# ==========================
# PROCESS PJ
# ==========================

pj = pd.read_excel(
    SOURCE_FILE,
    sheet_name="PJ"
)

for _, rec in pj.iterrows():

    date_found = False

    pj_dates = [
        "FR-CAN ID_Date",
        "SR-CAN ID_Date",
        "SR2-CAN ID_Date",
        "SR3-CAN ID_Date",
        "FR-EMP ID_Date",
        "SR-EMP ID_Date",
        "SR2-EMP ID_Date",
        "SR3-EMP ID_Date"
    ]

    for col in pj_dates:

        dt = to_dt(rec[col])

        if dt and dt.date() == dashboard_date.date():
            date_found = True
            break
    
    if not date_found:
        continue

    cid = rec["C_ID"]

    existing_row = find_existing_row(
        dst_ws,
        cid,
        rec.get("Case_SDate")
    )

    if existing_row:
        row = existing_row
    else:
        row = get_next_row(dst_ws)

        write_common(row, rec)

    latest_report_type = str(
        rec.get("Report Type", "")
    ).strip()

    if latest_report_type != "":
        dst_ws[f"H{row}"] = latest_report_type

    update_date_result_cells(
        row,
        rec,
        PJ_DATE_MAPPING
    )

    fr_can = pd.to_datetime(
        rec.get("FR-CAN ID_Date"),
        errors="coerce"
    )

    fr_emp = pd.to_datetime(
        rec.get("FR-EMP ID_Date"),
        errors="coerce"
    )
    
    if (
        pd.notna(fr_emp)
        and fr_emp.date() == dashboard_date.date()
        and pd.notna(fr_can)
        and fr_can.date() != dashboard_date.date()
    ):
        
        ce_status = str(
            rec.get("CE_Status", "")
        ).strip().upper()

        if ce_status in ["YES", "POB"]:

            ce1_result = str(
                rec.get("CE1_Result", "")
            ).strip().upper()

            if ce1_result in [
                "GREEN",
                "RED",
                "ORANGE",
                "MAROON"
            ]:
                dst_ws[f"Z{row}"] = 1

        # ReINT Additional Row

        reint = valid_result(
            rec["ReINT_GDB_Result"]
        )
    
        if (
            reint != "NA"
            and should_count_reint(
                rec,
                tracker_pj
            )
        ):
        
            row2 = get_next_row(dst_ws)

            write_common(row2, rec)
            update_date_result_cells(
                row2,
                rec,
                PJ_DATE_MAPPING
            )

            dst_ws[f"Y{row2}"] = "NA"
            dst_ws[f"Z{row2}"] = "NA"
            dst_ws[f"AA{row2}"] = "NA"
            dst_ws[f"AB{row2}"] = "NA"
            dst_ws[f"AC{row2}"] = "NA"
            dst_ws[f"AD{row2}"] = "NA"
            dst_ws[f"AE{row2}"] = "NA"
            dst_ws[f"AF{row2}"] = "NA"

            dst_ws[f"AG{row2}"] = reint

        continue

    sr_can = pd.to_datetime(
        rec.get("SR-CAN ID_Date"),
        errors="coerce"
    )

    sr2_can = pd.to_datetime(
        rec.get("SR2-CAN ID_Date"),
        errors="coerce"
    )

    sr3_can = pd.to_datetime(
        rec.get("SR3-CAN ID_Date"),
        errors="coerce"
    )

    sr_emp = pd.to_datetime(
        rec.get("SR-EMP ID_Date"),
        errors="coerce"
    )

    sr2_emp = pd.to_datetime(
        rec.get("SR2-EMP ID_Date"),
        errors="coerce"
    )

    sr3_emp = pd.to_datetime(
        rec.get("SR3-EMP ID_Date"),
        errors="coerce"
    )
    
    # FR-CAN / FR-EMP already delivered
    # Current SR-EMP / SR2-EMP / SR3-EMP
    # Do not count any checks

    if (
        (
            (pd.notna(sr_emp) and sr_emp.date() == dashboard_date.date())
            or
            (pd.notna(sr2_emp) and sr2_emp.date() == dashboard_date.date())
            or
            (pd.notna(sr3_emp) and sr3_emp.date() == dashboard_date.date())
        )
        and
        (
            pd.notna(fr_can)
            or pd.notna(fr_emp)
        )
    ):
        continue

    # FR-CAN / FR-EMP already delivered
    # Current SR-CAN / SR2-CAN / SR3-CAN
    # Do not count any checks

    if (
        (
            (pd.notna(sr_can) and sr_can.date() == dashboard_date.date())
            or
            (pd.notna(sr2_can) and sr2_can.date() == dashboard_date.date())
            or
            (pd.notna(sr3_can) and sr3_can.date() == dashboard_date.date())
        )
        and
        (
            pd.notna(fr_can)
            or pd.notna(fr_emp)
        )
    ):
        continue

    if (

        (
            (pd.notna(fr_emp) and fr_emp.date() == dashboard_date.date())
            or
            (pd.notna(sr_emp) and sr_emp.date() == dashboard_date.date())
            or
            (pd.notna(sr2_emp) and sr2_emp.date() == dashboard_date.date())
            or
            (pd.notna(sr3_emp) and sr3_emp.date() == dashboard_date.date())
        )

        and pd.isna(fr_can)
        and pd.isna(sr_can)
        and pd.isna(sr2_can)
        and pd.isna(sr3_can)

    ):

        emp_results = [
            rec.get("FR-EMP ID_Result"),
            rec.get("SR-EMP ID_Result"),
            rec.get("SR2-EMP ID_Result"),
            rec.get("SR3-EMP ID_Result")
        ]

        ce1_result = str(
            rec.get("CE1_Result", "")
        ).strip().upper()

        if ce1_result in [
            "GREEN",
            "RED",
            "ORANGE",
            "MAROON"
        ]:
            dst_ws[f"Z{row}"] = 1
              
      
    # Education

    if not is_duplicate_check(
        rec,
        tracker_pj,
        "Ed_SDate",
        "Ed_Result",
        "Ed_Edate"
    ):
        update_if_blank(
            dst_ws[f"Y{row}"],
            valid_result(rec["Ed_Result"])
        )

    # CE1

    ce_status = str(
        rec["CE_Status"]
    ).upper()

    ce1 = valid_result(
        rec["CE1_Result"]
    )

    if ce_status in [
        "YES",
        "POB"
    ]:
        update_if_blank(
            dst_ws[f"Z{row}"],
            "NA"
        )
    else:
        update_if_blank(
            dst_ws[f"Z{row}"],
            ce1
        )

    update_if_blank(
        dst_ws[f"AA{row}"],
        valid_result(rec["E2_Result"])
    )

    update_if_blank(
        dst_ws[f"AB{row}"],
        valid_result(rec["E3_Result"])
    )

    update_if_blank(
        dst_ws[f"AC{row}"],
        valid_result(rec["E4_Result"])
    )

    update_if_blank(
        dst_ws[f"AD{row}"],
        valid_result(rec["E5_Result"])
    )

    update_if_blank(
        dst_ws[f"AE{row}"],
        valid_result(rec["E6_Result"])
    )

    if not is_duplicate_check(
        rec,
        tracker_pj,
        "eP_CRC_SDate",
        "eP_CRC_Result",
        "eP_CRC_EDate"
    ):
        update_if_blank(
            dst_ws[f"AF{row}"],
            valid_result(rec["eP_CRC_Result"])
        )
    
    if not is_duplicate_check(
        rec,
        tracker_pj,
        "GDB_SDate",
        "GDB_Result",
        "GDB_EDate"
    ):
        update_if_blank(
            dst_ws[f"AG{row}"],
            valid_result(rec["GDB_Result"])
        )
    
# ==========================
# PROCESS PE & CVC
# ==========================

pe = pd.read_excel(
    SOURCE_FILE,
    sheet_name="PE & CVC"
)

for _, rec in pe.iterrows():
    
    date_found = False

    pe_dates = [
        "IR_Date",
        "FR_Date",
        "SR_Date",
        "SR2_Date",
        "SR3_Date"
    ]

    for col in pe_dates:

        dt = to_dt(rec[col])

        if dt and dt.date() == dashboard_date.date():
            date_found = True
            break

    if not date_found:
        continue

    cid = rec["C_ID"]

    existing_row = find_existing_row(
        dst_ws,
        cid,
        rec.get("Case_SDate")
    )

    if existing_row:
        row = existing_row
    else:
        row = get_next_row(dst_ws)

        write_common(row, rec)

    latest_report_type = str(
        rec.get("Report Type", "")
    ).strip()

    if latest_report_type != "":
        dst_ws[f"H{row}"] = latest_report_type

    update_date_result_cells(
        row,
        rec,
        PE_DATE_MAPPING
    )

    # PE OLD IR + CURRENT FR = CE1 ONLY

    ir_date = pd.to_datetime(
        rec.get("IR_Date"),
        errors="coerce"
    )

    fr_date = pd.to_datetime(
        rec.get("FR_Date"),
        errors="coerce"
    )

    if (
        pd.notna(fr_date)
        and fr_date.date() == dashboard_date.date()
        and pd.notna(ir_date)
        and ir_date.date() != dashboard_date.date()
    ):

        fr_result = str(
            rec.get("FR_Result", "")
        ).strip().upper()

        if fr_result in [
            "GREEN",
            "RED",
            "ORANGE",
            "MAROON"
        ]:
            dst_ws[f"Z{row}"] = 1

        continue
    
    sr_date = pd.to_datetime(
        rec.get("SR_Date"),
        errors="coerce"
    )

    sr2_date = pd.to_datetime(
        rec.get("SR2_Date"),
        errors="coerce"
    )

    sr3_date = pd.to_datetime(
        rec.get("SR3_Date"),
        errors="coerce"
    )

    # IR / FR already delivered
    # Current SR / SR2 / SR3
    # Details update only

    if (
        (
            (pd.notna(sr_date) and sr_date.date() == dashboard_date.date())
            or
            (pd.notna(sr2_date) and sr2_date.date() == dashboard_date.date())
            or
            (pd.notna(sr3_date) and sr3_date.date() == dashboard_date.date())
        )
        and
        (
            pd.notna(ir_date)
            or pd.notna(fr_date)
        )
    ):
        continue

    if (

        (
            (pd.notna(fr_date) and fr_date.date() == dashboard_date.date())
            or
            (pd.notna(sr_date) and sr_date.date() == dashboard_date.date())
            or
            (pd.notna(sr2_date) and sr2_date.date() == dashboard_date.date())
            or
            (pd.notna(sr3_date) and sr3_date.date() == dashboard_date.date())
        )

        and pd.isna(ir_date)

    ):

        pe_results = [
            rec.get("FR_Result"),
            rec.get("SR_Result"),
            rec.get("SR2_Result"),
            rec.get("SR3_Result")
        ]

        ce1_result = str(
            rec.get("CE1_Result", "")
        ).strip().upper()

        if ce1_result in [
            "GREEN",
            "RED",
            "ORANGE",
            "MAROON"
        ]:
            dst_ws[f"Z{row}"] = 1
        
    update_if_blank(
        dst_ws[f"Y{row}"],
        valid_result(rec["Ed_Result"])
    )     
    
    # IR only case -> CE1 ignore

    if (
        pd.notna(ir_date)
        and ir_date.date() == dashboard_date.date()
        and pd.isna(fr_date)
        and pd.isna(sr_date)
        and pd.isna(sr2_date)
        and pd.isna(sr3_date)
    ):
        pass

    else:

        ce_status = str(
            rec["CE_Status"]
        ).upper()

        ce1 = valid_result(
            rec["CE1_Result"]
        )

        if ce_status in ["YES", "POB"]:
            update_if_blank(
                dst_ws[f"Z{row}"],
                ce1
            )
        else:
            update_if_blank(
                dst_ws[f"Z{row}"],
                "NA"
            )

    update_if_blank(
        dst_ws[f"AA{row}"],
        valid_result(rec["E2_Result"])
    )

    update_if_blank(
        dst_ws[f"AB{row}"],
        valid_result(rec["E3_Result"])
    )

    update_if_blank(
        dst_ws[f"AC{row}"],
        valid_result(rec["E4_Result"])
    )

    update_if_blank(
        dst_ws[f"AD{row}"],
        valid_result(rec["E5_Result"])
    )

    update_if_blank(
        dst_ws[f"AF{row}"],
        valid_result(rec["eP_CRC_Result"])
    )

    update_if_blank(
        dst_ws[f"AG{row}"],
        valid_result(rec["GDB_Result"])
    )

# ==========================
# PROCESS FRESHER
# ==========================

fresher = pd.read_excel(
    SOURCE_FILE,
    sheet_name="Fresher"
)

for _, rec in fresher.iterrows():

    date_found = False
    
    fresher_dates = [
        "FR-CAN ID_Date",
        "SR-CAN ID_Date",
        "SR2-CAN ID_Date",
        "SR3-CAN ID_Date",
        "FR-EMP ID_Date",
        "SR-EMP ID_Date",
        "SR2-EMP ID_Date",
        "SR3-EMP ID_Date"
    ]

    for col in fresher_dates:

        dt = to_dt(rec[col])

        if dt and dt.date() == dashboard_date.date():
            date_found = True
            break

    if not date_found:
        continue

    cid = rec["C_ID"]

    existing_row = find_existing_row(
        dst_ws,
        cid,
        rec.get("Case_SDate")
    )

    if existing_row:
        row = existing_row
    else:
        row = get_next_row(dst_ws)

        write_common(row, rec)

    latest_report_type = str(
        rec.get("Report Type", "")
    ).strip()

    if latest_report_type != "":
        dst_ws[f"H{row}"] = latest_report_type

    update_date_result_cells(
        row,
        rec,
        PJ_DATE_MAPPING
    )

    # OLD FR-CAN + CURRENT FR-EMP

    fr_can = pd.to_datetime(
        rec.get("FR-CAN ID_Date"),
        errors="coerce"
    )

    fr_emp = pd.to_datetime(
        rec.get("FR-EMP ID_Date"),
        errors="coerce"
    )

    sr_can = pd.to_datetime(
        rec.get("SR-CAN ID_Date"),
        errors="coerce"
    )

    sr2_can = pd.to_datetime(
        rec.get("SR2-CAN ID_Date"),
        errors="coerce"
    )

    sr3_can = pd.to_datetime(
        rec.get("SR3-CAN ID_Date"),
        errors="coerce"
    )

    sr_emp = pd.to_datetime(
        rec.get("SR-EMP ID_Date"),
        errors="coerce"
    )

    sr2_emp = pd.to_datetime(
        rec.get("SR2-EMP ID_Date"),
        errors="coerce"
    )

    sr3_emp = pd.to_datetime(
        rec.get("SR3-EMP ID_Date"),
        errors="coerce"
    )

    # FR-CAN / FR-EMP already delivered
    # Current SR-EMP / SR2-EMP / SR3-EMP
    # Skip all counts

    history_exists = (
        pd.notna(fr_can)
        or pd.notna(fr_emp)
        or (
            pd.notna(sr_can)
            and sr_can.date() != dashboard_date.date()
        )
        or (
            pd.notna(sr2_can)
            and sr2_can.date() != dashboard_date.date()
        )
        or (
            pd.notna(sr3_can)
            and sr3_can.date() != dashboard_date.date()
        )
        or (
            pd.notna(sr_emp)
            and sr_emp.date() != dashboard_date.date()
        )
        or (
            pd.notna(sr2_emp)
            and sr2_emp.date() != dashboard_date.date()
        )
        or (
            pd.notna(sr3_emp)
            and sr3_emp.date() != dashboard_date.date()
        )
    )

    current_supp = (
        (
            pd.notna(sr_can)
            and sr_can.date() == dashboard_date.date()
        )
        or
        (
            pd.notna(sr2_can)
            and sr2_can.date() == dashboard_date.date()
        )
        or
        (
            pd.notna(sr3_can)
            and sr3_can.date() == dashboard_date.date()
        )
        or
        (
            pd.notna(sr_emp)
            and sr_emp.date() == dashboard_date.date()
        )
        or
        (
            pd.notna(sr2_emp)
            and sr2_emp.date() == dashboard_date.date()
        )
        or
        (
            pd.notna(sr3_emp)
            and sr3_emp.date() == dashboard_date.date()
        )
    )

    if current_supp and history_exists:
        continue

    # Existing logic
    if (
        pd.notna(fr_emp)
        and fr_emp.date() == dashboard_date.date()
        and pd.notna(fr_can)
        and fr_can.date() != dashboard_date.date()
    ):
        continue

        update_if_blank(
            dst_ws[f"Y{row}"],
            valid_result(rec["Ed_Result"])
        )

        update_if_blank(
            dst_ws[f"AF{row}"],
            valid_result(rec["eP_CRC_Result"])
        )

        update_if_blank(
            dst_ws[f"AG{row}"],
            valid_result(rec["GDB_Result"])
        )

        continue

    dst_ws[f"Z{row}"] = "NA"
    dst_ws[f"AA{row}"] = "NA"
    dst_ws[f"AB{row}"] = "NA"
    dst_ws[f"AC{row}"] = "NA"
    dst_ws[f"AD{row}"] = "NA"
    dst_ws[f"AE{row}"] = "NA"
                
    if not is_duplicate_check(
        rec,
        tracker_pj,
        "Ed_SDate",
        "Ed_Result",
        "Ed_Edate"
    ):
        update_if_blank(
            dst_ws[f"Y{row}"],
            valid_result(rec["Ed_Result"])
        )

    if not is_duplicate_check(
        rec,
        tracker_pj,
        "eP_CRC_SDate",
        "eP_CRC_Result",
        "eP_CRC_EDate"
    ):
        update_if_blank(
            dst_ws[f"AF{row}"],
            valid_result(rec["eP_CRC_Result"])
        )

    if not is_duplicate_check(
        rec,
        tracker_pj,
        "GDB_SDate",
        "GDB_Result",
        "GDB_EDate"
    ):
        update_if_blank(
            dst_ws[f"AG{row}"],
            valid_result(rec["GDB_Result"])
        )

# ==========================
# CONVERT RESULTS TO 1 / 0
# ==========================

result_columns = [
    "Y","Z","AA","AB","AC","AD","AE","AF","AG"
]

count_value_by_column = {
    "Y": 400,                 # Education
    "Z": 225, "AA": 225, "AB": 225, "AC": 225, "AD": 225, "AE": 225,
    "AF": 300,                 # Criminal
    "AG": 200                  # Global / GDC
}

for row in range(3, dst_ws.max_row + 1):

    if not dst_ws[f"B{row}"].value:
        continue

    for col in result_columns:

        cell = dst_ws[f"{col}{row}"]

        value = str(cell.value).strip().upper()

        if value in [
            "GREEN",
            "RED",
            "ORANGE",
            "MAROON"
        ]:

            cell.value = count_value_by_column[col]

        elif value in [
            "NA",
            "PENDING",
            "STOP CHECK",
            "",
            "-",
            "NONE"
        ]:

            cell.value = 0

        else:
            try:
                if float(cell.value) == 1:
                    cell.value = count_value_by_column[col]
            except (TypeError, ValueError):
                pass


# ==========================
# SAVE
# ==========================

dst_wb.save(TARGET_FILE)

print("\n════════════════════════════════════════════════════════════")
print("🔥 INFOSYS DATASHEET AUTOMATION ENGINE 🐦‍🔥")
print("════════════════════════════════════════════════════════════")

print(f"\033[96m📅 Dashboard Date : {dashboard_date_str}\033[0m")
print(f"\033[93m📄 Target File    : {TARGET_FILE}\033[0m")

print("\n\033[92m✅ EXECUTION SUCCESSFUL\033[0m")
print("\n\033[95m🚀 Datasheet Updated & Ready for Reporting 🚀\033[0m")

print("\n════════════════════════════════════════════════════════════")
print("\033[94m🏆 PROCESS COMPLETED SUCCESSFULLY 🏆\033[0m")
print("════════════════════════════════════════════════════════════")
