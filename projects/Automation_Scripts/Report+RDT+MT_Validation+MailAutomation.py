import pandas as pd
import os
import re
import win32com.client

from openpyxl import load_workbook
from datetime import datetime, timedelta

# =====================================================
# PATHS
# =====================================================

case_file = r"C:\Users\Balakumaran VR\Downloads\Case Validation.xlsx"

downloads_folder = r"C:\Users\Balakumaran VR\Downloads"

rdt_files = [
    os.path.join(downloads_folder, f)
    for f in os.listdir(downloads_folder)
    if f.startswith("Report Dispatch Tracker")
    and f.endswith(".xlsx")
]

rdt_file = max(
    rdt_files,
    key=os.path.getmtime
)

print(
    f"\nRDT File Used : {os.path.basename(rdt_file)}"
)

reports_folder = r"C:\Users\Balakumaran VR\Downloads\Reports"

# =====================================================
# CLEAN ID
# =====================================================

def clean_id(value):

    if pd.isna(value):
        return ""

    value = str(value).strip()

    if value.lower() in ["nan", "na", "-", ""]:
        return ""

    if value.endswith(".0"):
        value = value[:-2]

    return value

# =====================================================
# VALIDATION DATE SELECTION
# =====================================================

print("\n1 - Today")
print("2 - Yesterday")
print("3 - Custom Date(s)")

choice = input("\nChoose Option : ")

if choice == "1":

    validation_dates = [
        datetime.today().date()
    ]

elif choice == "2":

    validation_dates = [
        (datetime.today() - timedelta(days=1)).date()
    ]

elif choice == "3":

    user_input = input(
        "\nEnter Date(s) (DD-MMM-YYYY, comma separated) : "
    )

    validation_dates = [
        datetime.strptime(
            d.strip(),
            "%d-%b-%Y"
        ).date()
        for d in user_input.split(",")
    ]

else:

    print("Invalid Option")
    exit()

print(f"\nValidation Date(s) : {validation_dates}")

# =====================================================
# GET REPORT IDs FROM ALL SUBFOLDERS
# =====================================================

report_ids = set()
all_report_files = []
report_locations = {}

for root, dirs, files in os.walk(reports_folder):

    for file in files:

        if file.lower().endswith(".pdf"):

            all_report_files.append(file)

        match = re.match(r"^(\d+)_", file)

        if match:

            report_id = match.group(1)

            report_ids.add(report_id)

            report_locations[report_id] = root

# =====================================================
# LOAD RDT
# =====================================================

rdt = pd.read_excel(
    rdt_file,
    sheet_name="Data",
    engine="openpyxl"
)

wb = load_workbook(rdt_file)
ws = wb["Data"]

rdt_pairs = set()
rdt_report_types = {}
rdt_case_status_errors = []
rdt_hyphen_errors = []
rdt_row_blank_errors = []
rdt_bold_hyphen_errors = []

for excel_row, (_, row) in enumerate(
    rdt.iterrows(),
    start=2
):

    report_type = str(
        row["Report Type"]
    ).strip().upper()

    c_id = clean_id(row["C_ID"])
    emp_no = clean_id(row["Emp_No"])

    if c_id:
        rdt_pairs.add((c_id, report_type))

    if emp_no:
        rdt_pairs.add((emp_no, report_type))

    if c_id:
        rdt_report_types[c_id] = report_type
    
    if emp_no:
        rdt_report_types[emp_no] = report_type

    case_status = str(
        row["Case Status"]
    ).strip()

    has_candidate_data = (
        c_id != ""
        or emp_no != ""
        or str(row["C_Name"]).strip() not in ["", "nan"]
    )

    validation_id = c_id if c_id else emp_no

    if validation_id == "":

        validation_id = str(
            row.get("C_Name", "")
        ).strip()

        if validation_id.lower() == "nan":
            validation_id = ""

    if has_candidate_data:

        if case_status != "Report Dispatched to Client":

            rdt_case_status_errors.append(
                f"{validation_id} | Expected='Report Dispatched to Client' | Actual='{case_status}'"
            )

    else:

        if case_status == "Report Dispatched to Client":

            rdt_case_status_errors.append(
               "Case Status populated but candidate details missing"
            )

    mandatory_columns = [
        "Package Name",
        "Report Type",
        "IR / CID-Final & SR_Date",
        "IR / CID-Final & SR_Result",
        "EMP-Final_Date",
        "EMP-Final_Result",
        "EMP-SR_Date",
        "EMP-SR_Result",
        "EMP-SR2_Date",
        "EMP-SR2_Result",
        "EMP-SR3_Date",
        "EMP-SR3_Result",
        "Case Status"
    ]

    if has_candidate_data:

        for col in mandatory_columns:

            value = row.get(col)

            if pd.isna(value) or str(value).strip() == "":

                rdt_row_blank_errors.append(
                    f"{validation_id} | Blank Found in '{col}'"
                )

    result_columns = [
        "IR / CID-Final & SR_Date",
        "EMP-Final_Date",
        "EMP-SR_Date",
        "EMP-SR2_Date",
        "EMP-SR3_Date",
        "IR / CID-Final & SR_Result",
        "EMP-Final_Result",
        "EMP-SR_Result",
        "EMP-SR2_Result",
        "EMP-SR3_Result"
    ]

    for col in result_columns:

        try:

            col_no = rdt.columns.get_loc(col) + 1

            cell = ws.cell(
                row=excel_row,
                column=col_no
            )

            if (
                str(cell.value).strip() == "-"
                and cell.font.bold
            ):

                rdt_bold_hyphen_errors.append(
                    f"{validation_id} | {col} | Hyphen is Bold"
                )

        except:
            pass

# ==================================================
# HYPHEN VALIDATION
# ==================================================

    check_pairs = [
        ("EMP-Final_Date", "EMP-Final_Result"),
        ("EMP-SR_Date", "EMP-SR_Result"),
        ("EMP-SR2_Date", "EMP-SR2_Result"),
        ("EMP-SR3_Date", "EMP-SR3_Result")
    ]

    for date_col, result_col in check_pairs:

        date_value = row.get(date_col)
        result_value = row.get(result_col)

        date_blank = pd.isna(date_value) or str(date_value).strip() == ""
        result_blank = pd.isna(result_value) or str(result_value).strip() == ""

        # rendu field-um blank irukka koodadhu
        if date_blank and result_blank:

            rdt_hyphen_errors.append(
                f"{validation_id} | {date_col}/{result_col} | Expected '-'"
            )

        # onnu blank, onnu value irundhalum error
        elif date_blank != result_blank:

            rdt_hyphen_errors.append(
                f"{validation_id} | {date_col}/{result_col} | Date/Result mismatch"
            )

# =====================================================
# SHEET RULES
# =====================================================

sheet_configs = {

    "PJ": {
        "C_ID": [
            "FR-CAN ID_Date",
            "SR-CAN ID_Date",
            "SR2-CAN ID_Date",
            "SR3-CAN ID_Date"
        ],

        "Emp_No": [
            "FR-EMP ID_Date",
            "SR-EMP ID_Date",
            "SR2-EMP ID_Date",
            "SR3-EMP ID_Date"
        ]
    },

    "PE & CVC": {
        "C_ID": [
            "IR_Date"
        ],

        "Emp_No": [
            "FR_Date",
            "SR_Date",
            "SR2_Date",
            "SR3_Date"
        ]
    },

    "Fresher": {
        "C_ID": [
            "FR-CAN ID_Date",
            "SR-CAN ID_Date",
            "SR2-CAN ID_Date",
            "SR3-CAN ID_Date"
        ],

        "Emp_No": [
            "FR-EMP ID_Date",
            "SR-EMP ID_Date",
            "SR2-EMP ID_Date",
            "SR3-EMP ID_Date"
        ]
    }
}

# =====================================================
# DATE -> RESULT MAPPING
# =====================================================

result_mapping = {
    "FR-CAN ID_Date": "FR-CAN ID_Result",
    "SR-CAN ID_Date": "SR-CAN ID_Result",
    "SR2-CAN ID_Date": "SR2-CAN ID_Result",
    "SR3-CAN ID_Date": "SR3-CAN ID_Result",

    "FR-EMP ID_Date": "FR-EMP ID_Result",
    "SR-EMP ID_Date": "SR-EMP ID_Result",
    "SR2-EMP ID_Date": "SR2-EMP ID_Result",
    "SR3-EMP ID_Date": "SR3-EMP ID_Result",

    "IR_Date": "IR_Result",
    "FR_Date": "FR_Result",
    "SR_Date": "SR_Result",
    "SR2_Date": "SR2_Result",
    "SR3_Date": "SR3_Result"
}

# =====================================================
# RESULTS
# =====================================================

multiple_errors = []
missing_rdt = []
missing_reports = []
wrong_folder_errors = []
missing_validation_date = []
report_type_errors = []
rdt_report_type_errors = []
missing_color_code = []

# =====================================================
# PROCESS SHEETS
# =====================================================

for sheet_name, config in sheet_configs.items():

    print(f"\nChecking Sheet : {sheet_name}")

    df = pd.read_excel(
        case_file,
        sheet_name=sheet_name,
        engine="openpyxl"
    )

    for row_no, row in df.iterrows():

        matched = []
                
        # -------------------------
        # C_ID DATE COLUMNS
        # -------------------------

        for col in config["C_ID"]:

            value = row.get(col)

            if pd.notna(value):

                try:
                    if pd.to_datetime(value).date() in validation_dates:
                        matched.append(("C_ID", col))
                except:
                    pass

        # -------------------------
        # EMP_NO DATE COLUMNS
        # -------------------------

        for col in config["Emp_No"]:

            value = row.get(col)

            if pd.notna(value):

                try:
                    if pd.to_datetime(value).date() in validation_dates:
                        matched.append(("Emp_No", col))
                except:
                    pass

        # -------------------------
        # MULTIPLE COLUMN ERROR
        # -------------------------

        if len(matched) > 1:

            c_id = clean_id(row.get("C_ID"))
            emp_no = clean_id(row.get("Emp_No"))
            candidate_name = str(row.get("C_Name", "")).strip()

            multiple_errors.append(
                f"{sheet_name} | "
                f"Row {row_no + 2} | "
                f"Name={candidate_name} | "
                f"C_ID={c_id} | "
                f"Emp_No={emp_no} | "
                f"Columns={', '.join([x[1] for x in matched])}"
            )

            continue
        
        if len(matched) == 0:

            c_id = clean_id(row.get("C_ID"))
            emp_no = clean_id(row.get("Emp_No"))

            validation_id = c_id if c_id else emp_no

            if validation_id == "":

                validation_id = str(
                    row.get("C_Name", "")
                ).strip()

            report_type = str(
                row["Report Type"]
            ).strip().upper()

            rdt_found = (
                (c_id and (c_id, report_type) in rdt_pairs)
                or
                (emp_no and (emp_no, report_type) in rdt_pairs)
            )

            if (
                not rdt_found
                and validation_id != ""
                and c_id == ""
                and emp_no == ""
            ):

                candidate_name = str(
                    row.get("C_Name", "")
                ).strip().lower()

                rdt_found = any(
                    candidate_name == str(name).strip().lower()
                    for name in rdt["C_Name"].fillna("")
                )

            report_found = validation_id in report_ids

            same_id_report_found = False

            same_rows = df[
                (
                    (str(c_id).strip() != "")
                    &
                    (
                        df["C_ID"].astype(str).str.strip()
                        ==
                        str(c_id).strip()
                    )
                )
                |
                (
                    (str(emp_no).strip() != "")
                    &
                    (
                        df["Emp_No"].astype(str).str.strip()
                        ==
                        str(emp_no).strip()
                    )
                )
            ]

            for _, dup_row in same_rows.iterrows():

                for col in (
                    config["C_ID"] +
                    config["Emp_No"]
                ):

                    try:

                        val = dup_row.get(col)

                        if (
                            pd.notna(val)
                            and pd.to_datetime(val).date()
                            in validation_dates
                        ):
                            same_id_report_found = True
                            break

                    except:
                        pass

                if same_id_report_found:
                    break

            if same_id_report_found:
                continue


            if rdt_found or report_found:

                message = (
                    f"{sheet_name} | {validation_id} | No Date Available"
                )

                if message not in missing_validation_date:
                    missing_validation_date.append(message)

            continue
        
        selected_id_type = matched[0][0]

        validation_id = clean_id(
            row[selected_id_type]
        )

        if validation_id == "":
            continue

        report_type = str(
            row["Report Type"]
        ).strip().upper()

        package_name = str(
            row.get("Package Name", "")
        ).strip()

        expected_folder = None

        if package_name in [
            "Pre-Joining",
            "Pre-joining Lateral",
            "Post Employment"
        ]:
            expected_folder = "Lateral"

        elif package_name in [
            "CVC Lateral",
            "CVC Fresher"
        ]:
            expected_folder = "CVC"

        elif package_name == "Fresher":
            expected_folder = "Fresher"

        if report_type == "" or report_type == "NAN":

            report_type_errors.append(
                f"{sheet_name} | {validation_id} | REPORT TYPE BLANK"
            )

            continue

        expected_report_type = None

        stop_check_columns = [
            "FR-CAN ID_Result",
            "FR-EMP ID_Result",
            "FR_Result"
        ]

        for col in stop_check_columns:

            if col in df.columns:

                value = str(row.get(col)).strip().upper()

                if value == "STOP CHECK":

                    expected_report_type = "STOP CHECK"
                    break
        
        # PE & CVC

        if sheet_name == "PE & CVC" and expected_report_type != "STOP CHECK":
      
            if pd.notna(row.get("SR3_Date")):
                expected_report_type = "SUPPLEMENTARY-3"

            elif pd.notna(row.get("SR2_Date")):
                expected_report_type = "SUPPLEMENTARY-2"

            elif pd.notna(row.get("SR_Date")):
                expected_report_type = "SUPPLEMENTARY"

            elif pd.notna(row.get("FR_Date")):
                expected_report_type = "FINAL"

            elif pd.notna(row.get("IR_Date")):
                expected_report_type = "INTERIM"
        
        # CAN ID Dates

        if str(row.get("SR3-CAN ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "SUPPLEMENTARY-3"

        elif str(row.get("SR2-CAN ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "SUPPLEMENTARY-2"

        elif str(row.get("SR-CAN ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "SUPPLEMENTARY"

        elif str(row.get("FR-CAN ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "FINAL"

        
        # EMP ID Dates

        if str(row.get("SR3-EMP ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "SUPPLEMENTARY-3"

        elif str(row.get("SR2-EMP ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "SUPPLEMENTARY-2"

        elif str(row.get("SR-EMP ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
            expected_report_type = "SUPPLEMENTARY"

        elif str(row.get("FR-EMP ID_Date")).strip() not in ["", "nan", "NaT", "None"]:
           expected_report_type = "FINAL"

                            
        if expected_report_type and report_type != expected_report_type:

            report_type_errors.append(
                f"{sheet_name} | {validation_id} | Expected={expected_report_type} | Actual={report_type}"
            )


        if validation_id in rdt_report_types:

            rdt_report_type = rdt_report_types[validation_id]

            if expected_report_type and rdt_report_type != expected_report_type:

                rdt_report_type_errors.append(
                    f"{sheet_name} | {validation_id} | "
                    f"Expected={expected_report_type} | RDT={rdt_report_type}"
                )

        # =====================================================
        # COLOR CODE VALIDATION
        # =====================================================

        for date_col, result_col in result_mapping.items():

            if date_col in df.columns and result_col in df.columns:

                date_value = row.get(date_col)
                result_value = row.get(result_col)

                try:

                    if (
                        pd.notna(date_value)
                        and pd.to_datetime(date_value).date() in validation_dates
                    ):

                        if (
                            pd.isna(result_value)
                            or str(result_value).strip() == ""
                        ):

                            missing_color_code.append(
                                f"{sheet_name} | {validation_id} | {result_col} NOT FOUND"
                            )

                except:
                    pass
                    
        date_cols = config["C_ID"]

        for i in range(1, len(date_cols)):

            current_date = row.get(date_cols[i])
            previous_date = row.get(date_cols[i - 1])

            try:
                if (
                    pd.notna(current_date)
                    and pd.to_datetime(current_date).date() in validation_dates
                    and pd.isna(previous_date)
                ):

                    validation_id = clean_id(row.get("C_ID"))

                    missing_validation_date.append(
                        f"{sheet_name} | {validation_id} | {date_cols[i-1]} NOT FOUND"
                    )

            except:
                pass

        # -------------------------
        # RDT CHECK
        # -------------------------

        if validation_id not in rdt_report_types:

            missing_rdt.append(
                f"{sheet_name} | {validation_id} | {report_type}"
            )

        # -------------------------
        # REPORT CHECK
        # -------------------------

        if validation_id not in report_ids:

            print(f"Missing Report -> {validation_id}")

            missing_reports.append(
                f"{sheet_name} | {validation_id} | {report_type}"
            )

        actual_folder = report_locations.get(validation_id)

        if actual_folder:

            if "\\Lateral\\" in actual_folder:
                actual_folder = "Lateral"

            elif "\\CVC\\" in actual_folder:
                actual_folder = "CVC"

            elif "\\Fresher\\" in actual_folder:
                actual_folder = "Fresher"

        if (
            expected_folder
            and actual_folder
            and actual_folder != expected_folder
        ):

            wrong_folder_errors.append(
                f"{validation_id} | Package={package_name} | Expected={expected_folder} | Found={actual_folder}"
            )

# =====================================================
# OUTPUT
# =====================================================

print("\n")
print("=" * 80)
print("MULTIPLE DATE ERRORS")
print("=" * 80)

if multiple_errors:
    for item in multiple_errors:
        print(item)
else:
    print("No Multiple Date Errors")

print("\n")
print("=" * 80)
print("MISSING VALIDATION DATE")
print("=" * 80)

if missing_validation_date:
    for item in missing_validation_date:
        print(item)
else:
    print("No Missing Validation Date Records")

print("\n")
print("=" * 80)
print("MISSING COLOR CODE")
print("=" * 80)

if missing_color_code:
    for item in missing_color_code:
        print(item)
else:
    print("No Missing Color Code Records")

print("\n")
print("=" * 80)
print("REPORT TYPE ERRORS")
print("=" * 80)

if report_type_errors:
    for item in report_type_errors:
        print(item)
else:
    print("No Report Type Errors")

print("\n")
print("=" * 80)
print("RDT REPORT TYPE ERRORS")
print("=" * 80)

if rdt_report_type_errors:
    for item in rdt_report_type_errors:
        print(item)
else:
    print("No RDT Report Type Errors")

print("\n")
print("=" * 80)
print("RDT BOLD HYPHEN ERRORS")
print("=" * 80)

if rdt_bold_hyphen_errors:

    for item in rdt_bold_hyphen_errors:
        print(item)
else:
    print("No RDT Bold Hyphen Errors")

print("\n")
print("=" * 80)
print("RDT CASE STATUS ERRORS")
print("=" * 80)

if rdt_case_status_errors:
    for item in rdt_case_status_errors:
        print(item)
else:
    print("No RDT Case Status Errors")

print("\n")
print("=" * 80)
print("RDT HYPHEN ERRORS")
print("=" * 80)

if rdt_hyphen_errors:
    for item in rdt_hyphen_errors:
        print(item)
else:
    print("No RDT Hyphen Errors")

print("\n")
print("=" * 80)
print("RDT ROW BLANK ERRORS")
print("=" * 80)

if rdt_row_blank_errors:
    for item in rdt_row_blank_errors:
        print(item)
else:
    print("No RDT Row Blank Errors")

print("\n")
print("=" * 80)
print("MISSING IN RDT")
print("=" * 80)

if missing_rdt:
    for item in missing_rdt:
        print(item)
else:
    print("No Missing Records in RDT")

print("\n")
print("=" * 80)
print("WRONG REPORT FOLDER")
print("=" * 80)

if wrong_folder_errors:

    for item in wrong_folder_errors:
        print(item)

else:

    print("No Wrong Folder Errors")

print("\n")
print("=" * 80)
print("MISSING IN REPORTS")
print("=" * 80)

if missing_reports:
    for item in missing_reports:
        print(item)
else:
    print("No Missing Records in Reports")

print("\nValidation Completed Successfully")

# =====================================================
# MAIL AUTOMATION
# =====================================================

import html
import uuid
from pathlib import Path


def prompt_report_date():
    while True:
        date_text = input("Subject date (DD-MMM-YYYY, e.g. 07-Oct-2026): ").strip()
        try:
            return datetime.strptime(date_text, "%d-%b-%Y").strftime("%d-%b-%Y")
        except ValueError:
            print("Invalid date. Enter it as DD-MMM-YYYY, for example 07-Oct-2026.")


def get_recipient_and_subject(pdf_path, report_date):
    report_name = pdf_path.stem
    folder_path = str(pdf_path.parent).casefold()
    if "fresher" in folder_path:
        return "freshersbgc@infosys.com", f"{report_name} - {report_date}"
    if "cvc" in folder_path:
        return "lateralbgc@infosys.com", f"{report_name}_CVC - {report_date}"
    return "lateralbgc@infosys.com", f"{report_name} - {report_date}"


def build_html_body(note):
    note_text = (note or "").replace("\xa0", " ").replace("\u200b", "").replace("\ufeff", "").strip()
    body = (
        '<p style="margin:0 0 12pt 0">Dear Team,</p>'
        '<p style="margin:0 0 12pt 0">Please find the enclosed verified report '
        'with reference to the subject above.</p>'
    )
    if note_text:
        note_html = html.escape(note_text).replace("\r\n", "\n").replace("\n", "<br>")
        body += '<p style="margin:0 0 12pt 0"><strong>Note : ' + note_html + '</strong></p>'
    body += '<p style="margin:0">Regards,<br>BGV Team - DigiVerifier</p>'
    return '<html><body><div style="font-family:Calibri, sans-serif; font-size:11pt">' + body + '</div></body></html>'


print("\n")
print("=" * 60)
print("MAIL SUMMARY")
print("=" * 60)
print(f"Multiple Date Errors : {len(multiple_errors)}")
print(f"Missing In RDT       : {len(missing_rdt)}")
print(f"Report Type Errors   : {len(report_type_errors)}")
print(f"RDT Report Errors    : {len(rdt_report_type_errors)}")
print(f"Missing Validation   : {len(missing_validation_date)}")
print(f"Missing Reports      : {len(missing_reports)}")

choice = input("\nDo you want to create draft emails for VALID reports only? (Y/N): ").strip().upper()
if choice != "Y":
    print("\nOperation Cancelled By User")
    input("\nPress Enter to Exit...")
    exit()

# Exclude report IDs flagged by the validation results above.
exclude_ids = set()
for error_list in (
    missing_rdt,
    missing_reports,
    missing_validation_date,
    report_type_errors,
    rdt_report_type_errors,
):
    for item in error_list:
        parts = item.split("|")
        if len(parts) >= 2:
            exclude_ids.add(parts[1].strip())
for item in multiple_errors:
    match = re.search(r"C_ID=(\d+)", item)
    if match:
        exclude_ids.add(match.group(1))

print(f"\nExcluded IDs : {len(exclude_ids)}")

reports_path = Path(reports_folder)
if not reports_path.is_dir():
    print(f"Reports folder not found: {reports_path}")
    input("\nPress Enter to Exit...")
    exit()

pdf_files = sorted(
    (path for path in reports_path.rglob("*") if path.is_file() and path.suffix.casefold() == ".pdf"),
    key=lambda path: (str(path.parent).casefold(), path.name.casefold()),
)
valid_pdf_files = []
for pdf_path in pdf_files:
    report_id_match = re.match(r"^(\d+)_", pdf_path.name)
    if report_id_match and report_id_match.group(1) in exclude_ids:
        continue
    if not report_id_match:
        print(f"NON-STANDARD REPORT FILE (included): {pdf_path.name}")
    valid_pdf_files.append(pdf_path)

if not valid_pdf_files:
    print(f"No valid PDF reports found under: {reports_path}")
    input("\nPress Enter to Exit...")
    exit()

report_date = prompt_report_date()
note = input("Note to include in every email (leave blank for no Note): ").strip()
print(f"\nFound {len(valid_pdf_files)} valid PDF report(s). Subject date: {report_date}")
if input("Create Outlook drafts for these reports? (Y/N): ").strip().upper() != "Y":
    print("Cancelled. No drafts were created.")
    input("\nPress Enter to Exit...")
    exit()

SENDER_EMAIL = "infosys.bgvdoc@digiverifier.com"
outlook = win32com.client.Dispatch("Outlook.Application")
sender_account = None
for account in outlook.Session.Accounts:
    try:
        if str(account.SmtpAddress).casefold() == SENDER_EMAIL.casefold():
            sender_account = account
            break
    except Exception:
        continue
if sender_account is None:
    print(f"Outlook sender account not found: {SENDER_EMAIL}")
    input("\nPress Enter to Exit...")
    exit()

try:
    sender_drafts = sender_account.DeliveryStore.GetDefaultFolder(16)  # olFolderDrafts
except Exception as exc:
    print(f"Could not open the {SENDER_EMAIL} Drafts folder: {exc}")
    input("\nPress Enter to Exit...")
    exit()
print(f"Drafts will be saved to: {sender_drafts.FolderPath}")

mail_body = build_html_body(note)
cleanup_tag = "ReportMailCleanup_" + uuid.uuid4().hex
created_mails = []
for pdf_path in valid_pdf_files:
    recipient, subject = get_recipient_and_subject(pdf_path, report_date)
    try:
        mail = sender_drafts.Items.Add("IPM.Note")
        mail.SendUsingAccount = sender_account
        mail.To = recipient
        mail.Subject = subject
        mail.HTMLBody = mail_body
        mail.Attachments.Add(str(pdf_path))
        mail.BillingInformation = cleanup_tag
        mail.Save()
        created_mails.append(mail)
        print(f"Draft Created : {len(created_mails)} | {subject} | {pdf_path.name}")
    except Exception as exc:
        print(f"Draft Failed : {pdf_path.name} | {exc}")

print("\n" + "=" * 60)
print(f"TOTAL DRAFTS CREATED : {len(created_mails)} / {len(valid_pdf_files)}")
print("=" * 60)
if not created_mails:
    input("\nPress Enter to Exit...")
    exit()

if input("\nSend all created drafts now? (Y/N): ").strip().upper() == "Y":
    sent_count = 0
    for mail in created_mails:
        try:
            mail.BillingInformation = ""
            mail.Save()
            mail.Send()
            sent_count += 1
        except Exception as exc:
            print(f"Send Failed : {mail.Subject} | {exc}")
    print(f"Successfully sent: {sent_count} / {len(created_mails)}")
else:
    print("Send declined. Permanently deleting the drafts created in this run...")
    deleted_count = 0
    try:
        deleted_folder = sender_account.DeliveryStore.GetDefaultFolder(3)  # olFolderDeletedItems
        for mail in created_mails:
            try:
                mail.Delete()
            except Exception as exc:
                print(f"Draft removal failed : {mail.Subject} | {exc}")
        tagged_items = deleted_folder.Items.Restrict(f"[BillingInformation] = '{cleanup_tag}'")
        for item_index in range(tagged_items.Count, 0, -1):
            try:
                tagged_items.Item(item_index).Delete()
                deleted_count += 1
            except Exception as exc:
                print(f"Permanent deletion failed for a tagged draft: {exc}")
                break
        print(f"Permanently deleted from Outlook folders: {deleted_count} / {len(created_mails)}")
    except Exception as exc:
        print(f"Could not complete permanent draft cleanup: {exc}")
        print("Check the Infosys BGVDoc Deleted Items folder for this run's drafts.")

input("\nPress Enter to Exit...")
