import glob
import os
import pandas as pd
from openpyxl import load_workbook

# =====================================
# FILES
# =====================================

invoice_file = (
    r"C:\Users\Balakumaran VR\OneDrive - digiverifier solutions pvt ltd"
    r"\Shortcuts\Client Account Managers - Documents\General"
    r"\Infosys Conventional-AM-Invoice Sheet.xlsx"
)

downloads_folder = r"C:\Users\Balakumaran VR\Downloads"
datasheet_files = glob.glob(
    os.path.join(downloads_folder, "Infosys_Conventional-Datasheet*.xlsx")
)

if not datasheet_files:
    raise FileNotFoundError(
        "No Infosys_Conventional-Datasheet*.xlsx file found in Downloads."
    )

datasheet_file = max(
    datasheet_files,
    key=os.path.getmtime
)
print(f"Datasheet file used: {os.path.basename(datasheet_file)}")

# =====================================
# LIVE DIT COUNTS AND AMOUNTS
# =====================================

wb = load_workbook(
    invoice_file,
    data_only=True
)

ws = wb.active

def number_or_zero(value):
    if value is None or str(value).strip() == "":
        return 0.0
    try:
        return float(str(value).replace(",", "").replace("₹", "").strip())
    except (TypeError, ValueError):
        return 0.0

# Live DIT count cells (existing count cross-check)
live_dit_emp_count = number_or_zero(ws["H34"].value)
live_dit_edu_count = number_or_zero(ws["J34"].value)
live_dit_gdc_count = number_or_zero(ws["L34"].value)
live_dit_crc_count = number_or_zero(ws["N34"].value)

# Live DIT amount cells (new amount cross-check)
live_dit_emp_amount = number_or_zero(ws["I34"].value)
live_dit_edu_amount = number_or_zero(ws["K34"].value)
live_dit_gdc_amount = number_or_zero(ws["M34"].value)
live_dit_crc_amount = number_or_zero(ws["O34"].value)

# =====================================
# DATASHEET COUNTS AND AMOUNTS
# =====================================

df = pd.read_excel(
    datasheet_file,
    engine="openpyxl"
)

# Datasheet cells now hold per-check amounts rather than the old value 1.
# Accept both 1 and the configured amount so historical rows still count.
CHECK_AMOUNT = {
    "education": 400,
    "employment": 225,
    "crc": 300,
    "gdc": 200
}

def numeric_column(col_idx):
    return pd.to_numeric(df.iloc[:, col_idx], errors="coerce").fillna(0)

def check_count(col_idx, unit_amount):
    values = numeric_column(col_idx)
    return int(((values == 1) | (values == unit_amount)).sum())

def amount_sum(col_idx, unit_amount):
    values = numeric_column(col_idx)
    # Treat legacy 1s as the current per-check amount.
    values = values.where(values != 1, unit_amount)
    return float(values.sum())

def datasheet_totals(col_indices, unit_amount):
    count = sum(check_count(col_idx, unit_amount) for col_idx in col_indices)
    amount = sum(amount_sum(col_idx, unit_amount) for col_idx in col_indices)
    return count, amount

# Y = Education; Z:AE = Employment; AF = CRC; AG = GDC
datasheet_edu_count, datasheet_edu_amount = datasheet_totals(
    [24], CHECK_AMOUNT["education"]
)
datasheet_emp_count, datasheet_emp_amount = datasheet_totals(
    [25, 26, 27, 28, 29, 30], CHECK_AMOUNT["employment"]
)
datasheet_crc_count, datasheet_crc_amount = datasheet_totals(
    [31], CHECK_AMOUNT["crc"]
)
datasheet_gdc_count, datasheet_gdc_amount = datasheet_totals(
    [32], CHECK_AMOUNT["gdc"]
)

# =====================================
# VALIDATION
# =====================================

print("\n" + "=" * 70)
print("LIVE DIT vs DATASHEET COUNT VALIDATION")
print("=" * 70)

count_checks = [
    ("Education", live_dit_edu_count, datasheet_edu_count),
    ("Employment", live_dit_emp_count, datasheet_emp_count),
    ("GDC", live_dit_gdc_count, datasheet_gdc_count),
    ("CRC", live_dit_crc_count, datasheet_crc_count)
]

for name, live_dit_value, datasheet_value in count_checks:
    if live_dit_value == datasheet_value:
        status = "✅"
    else:
        status = "❌"
    print(
        f"{status} {name:<12} | "
        f"Live DIT Count={live_dit_value:g} | "
        f"Datasheet Count={datasheet_value}"
    )

print("\n" + "=" * 70)
print("LIVE DIT vs DATASHEET AMOUNT VALIDATION")
print("=" * 70)

amount_checks = [
    ("Education", live_dit_edu_amount, datasheet_edu_amount),
    ("Employment", live_dit_emp_amount, datasheet_emp_amount),
    ("GDC", live_dit_gdc_amount, datasheet_gdc_amount),
    ("CRC", live_dit_crc_amount, datasheet_crc_amount)
]

for name, live_dit_value, datasheet_value in amount_checks:
    if live_dit_value == datasheet_value:
        status = "✅"
    else:
        status = "❌"
    difference = abs(live_dit_value - datasheet_value)
    print(
        f"{status} {name:<12} | "
        f"Live DIT Amount={live_dit_value:g} | "
        f"Datasheet Amount={datasheet_value:g} | "
        f"Difference={difference:g}"
    )

print("=" * 70)